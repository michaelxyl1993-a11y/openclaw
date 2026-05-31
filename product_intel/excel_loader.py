"""Excel loader for Product Intel candidate products."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .input_adapter import normalize_product_rows


def clean_cell(value: Any) -> Any:
    if isinstance(value, str):
        return value.strip()
    return value


def row_has_value(values: list[Any]) -> bool:
    return any(str(value or "").strip() for value in values)


def trim_empty_columns(rows: list[list[Any]]) -> list[list[Any]]:
    if not rows:
        return rows
    max_len = max(len(row) for row in rows)
    used_indexes: list[int] = []
    for index in range(max_len):
        if any(index < len(row) and str(row[index] or "").strip() for row in rows):
            used_indexes.append(index)
    return [[row[index] if index < len(row) else "" for index in used_indexes] for row in rows]


def first_non_empty_sheet(workbook: Any) -> Any:
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            if row_has_value([clean_cell(value) for value in row]):
                return sheet
    return workbook.worksheets[0]


def load_raw_excel_rows(path: str | Path) -> tuple[list[dict[str, Any]], str]:
    excel_path = Path(path)
    file = excel_path.open("rb")
    try:
        workbook = load_workbook(file, read_only=True, data_only=True)
        sheet = first_non_empty_sheet(workbook)
        raw_rows: list[list[Any]] = []
        for row in sheet.iter_rows(values_only=True):
            values = [clean_cell(value) for value in row]
            if row_has_value(values):
                raw_rows.append(values)
        raw_rows = trim_empty_columns(raw_rows)
        if not raw_rows:
            return [], sheet.title

        headers = [str(value or "").strip() for value in raw_rows[0]]
        rows: list[dict[str, Any]] = []
        for raw_row in raw_rows[1:]:
            item = {header: raw_row[index] if index < len(raw_row) else "" for index, header in enumerate(headers) if header}
            if any(str(value or "").strip() for value in item.values()):
                rows.append(item)
        return rows, sheet.title
    finally:
        file.close()


def load_products_from_excel(path: str | Path, source: str = "auto") -> list[dict[str, Any]]:
    rows, _sheet_name = load_raw_excel_rows(path)
    return normalize_product_rows(rows, source=source)
